"""
DOCX document reader.
"""

from pathlib import Path

from docx import Document as DocxDocument

from minuta_generator.models import Document, DocumentFormat
from minuta_generator.readers.base import DocumentReader


class DOCXReader(DocumentReader):
	"""Reader for DOCX documents."""

	def supported_format(self) -> DocumentFormat:
		return DocumentFormat.DOCX

	def read(self, file_path: Path) -> Document:
		"""
		Read a DOCX document.

		Args:
			file_path: Path to the DOCX file

		Returns:
			Document object with extracted text content
		"""
		self.validate_file(file_path)

		docx_doc = DocxDocument(file_path)

		# Extract text from paragraphs
		paragraphs = []
		for para in docx_doc.paragraphs:
			if para.text.strip():
				paragraphs.append(para.text)

		# Extract text from tables
		table_texts = []
		for table in docx_doc.tables:
			for row in table.rows:
				row_text = ' | '.join(cell.text for cell in row.cells)
				if row_text.strip():
					table_texts.append(row_text)

		# Combine all content
		full_content = '\n'.join(paragraphs)
		if table_texts:
			full_content += '\n\n--- Tables ---\n' + '\n'.join(table_texts)

		# Extract metadata
		metadata = {
			'num_paragraphs': len(docx_doc.paragraphs),
			'num_tables': len(docx_doc.tables),
			'num_sections': len(docx_doc.sections),
		}

		# Try to extract core properties
		try:
			core_props = docx_doc.core_properties
			metadata.update(
				{
					'title': core_props.title or '',
					'author': core_props.author or '',
					'subject': core_props.subject or '',
					'keywords': core_props.keywords or '',
					'created': str(core_props.created) if core_props.created else '',
					'modified': str(core_props.modified) if core_props.modified else '',
				}
			)
		except Exception:
			pass  # Skip if metadata not available

		return Document(path=file_path, format=DocumentFormat.DOCX, content=full_content, metadata=metadata)
