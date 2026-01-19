"""
DOCX template writer.
"""

from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

from minuta_generator.models import Template


class DOCXWriter:
	"""Writes templates to DOCX format."""

	def write(self, template: Template, output_path: Path) -> None:
		"""
		Write template to DOCX file.

		Args:
			template: Template to write
			output_path: Path where DOCX file should be saved
		"""
		doc = self._generate_docx(template)
		doc.save(output_path)

	def _generate_docx(self, template: Template) -> DocxDocument:
		"""
		Generate DOCX document from template.

		Args:
			template: Template to convert

		Returns:
			DocxDocument object
		"""
		doc = DocxDocument()

		# Add title
		title = doc.add_heading(f'Template: {template.name}', level=1)

		# Add template content
		content_para = doc.add_paragraph()

		for section in template.sections:
			run = content_para.add_run(section.content)

			if section.is_variable:
				# Highlight variable sections in yellow
				run.font.highlight_color = WD_COLOR_INDEX.YELLOW
				run.bold = True
				run.font.color.rgb = RGBColor(133, 100, 4)  # Dark yellow/brown

				# Add comment-like notation
				if section.placeholder_name:
					comment_run = content_para.add_run(f' [{section.placeholder_name}]')
					comment_run.font.size = Pt(8)
					comment_run.font.color.rgb = RGBColor(128, 128, 128)
					comment_run.italic = True

		# Add separator
		doc.add_paragraph('─' * 80)

		# Add placeholders guide
		doc.add_heading('Placeholders Identificados', level=2)

		seen_placeholders = set()
		placeholder_items = []

		for section in template.sections:
			if section.is_variable and section.placeholder_name and section.placeholder_name not in seen_placeholders:
				seen_placeholders.add(section.placeholder_name)
				desc = section.description or section.pattern_type or 'N/A'
				placeholder_items.append(f'{section.placeholder_name}: {desc}')

		if placeholder_items:
			for item in placeholder_items:
				para = doc.add_paragraph(item, style='List Bullet')
		else:
			doc.add_paragraph('Nenhum placeholder identificado.')

		# Add metadata section
		doc.add_heading('Metadados', level=2)

		metadata_text = [
			f'Documentos originais: {len(template.original_documents)}',
			f'Padrões identificados: {template.metadata.get("num_patterns", 0)}',
		]

		if template.metadata.get('common_structure'):
			metadata_text.append(f'\nEstrutura comum:\n{template.metadata["common_structure"]}')

		for meta_item in metadata_text:
			doc.add_paragraph(meta_item)

		# Add instructions
		doc.add_page_break()
		doc.add_heading('Instruções de Uso', level=2)

		instructions = [
			'Este é um template gerado automaticamente a partir da análise de documentos jurídicos.',
			'',
			'Como usar:',
			'1. Textos destacados em amarelo são campos variáveis (placeholders)',
			'2. Substitua cada placeholder pelo valor apropriado para seu caso',
			'3. Os placeholders seguem o formato {{NOME_DO_CAMPO}}',
			'4. Mantenha o texto não destacado como está (conteúdo fixo)',
			'',
			'Consulte a lista de placeholders acima para entender o que cada campo representa.',
		]

		for instruction in instructions:
			if instruction:
				doc.add_paragraph(instruction)
			else:
				doc.add_paragraph()  # Empty line

		return doc
